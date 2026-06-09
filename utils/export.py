import os
from datetime import date
from typing import List

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from scrapers.base import RankItem, PLATFORMS, RANK_TYPES

# 平台展示顺序
PLATFORM_ORDER = ["appstore", "taptap", "bilibili", "huawei", "xiaomi", "oppo", "vivo", "honor", "kuaibao"]
RANK_TYPE_ORDER = ["active", "revenue", "download", "new"]

HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT  = Font(bold=True, color="FFFFFF", size=11)
ALT_FILL     = PatternFill("solid", fgColor="D6E4F0")
CENTER       = Alignment(horizontal="center", vertical="center")
THIN_BORDER  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

COL_HEADERS = ["排名", "游戏名称", "开发商", "评分", "游戏ID"]
COL_WIDTHS  = [8, 35, 25, 8, 20]


class Exporter:
    def __init__(self, output_dir: str = "./data"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def to_excel(self, items: List[RankItem], fetch_date: str = None) -> str:
        if not fetch_date:
            fetch_date = date.today().isoformat()

        # 按平台+榜单分组
        groups: dict = {}
        for item in items:
            key = (item.platform, item.rank_type)
            groups.setdefault(key, []).append(item)

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # 删除默认空 sheet

        # 每个平台独立一个 sheet
        for platform in PLATFORM_ORDER:
            sheet_data = {}
            for rank_type in RANK_TYPE_ORDER:
                key = (platform, rank_type)
                if key in groups:
                    sheet_data[rank_type] = groups[key]
            if not sheet_data:
                continue

            pname = PLATFORMS.get(platform, platform)
            ws = wb.create_sheet(title=pname)
            self._write_platform_sheet(ws, pname, sheet_data, fetch_date)

        # 汇总 sheet（总览）
        self._write_summary_sheet(wb, groups, fetch_date)

        path = os.path.join(self.output_dir, f"game_rank_{fetch_date}.xlsx")
        wb.save(path)
        return path

    def _write_platform_sheet(self, ws, platform_name: str, data: dict, fetch_date: str):
        row = 1
        for rank_type in RANK_TYPE_ORDER:
            items = data.get(rank_type)
            if not items:
                continue
            type_name = RANK_TYPES.get(rank_type, rank_type)

            # 榜单标题行
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COL_HEADERS))
            cell = ws.cell(row=row, column=1, value=f"{platform_name} · {type_name}（{fetch_date}）")
            cell.font = Font(bold=True, size=12, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E75B6")
            cell.alignment = CENTER
            row += 1

            # 列头
            for col, (header, width) in enumerate(zip(COL_HEADERS, COL_WIDTHS), 1):
                c = ws.cell(row=row, column=col, value=header)
                c.font = HEADER_FONT
                c.fill = HEADER_FILL
                c.alignment = CENTER
                c.border = THIN_BORDER
                ws.column_dimensions[get_column_letter(col)].width = width
            row += 1

            # 数据行
            for i, item in enumerate(items):
                fill = ALT_FILL if i % 2 == 0 else None
                vals = [item.rank, item.name, item.developer,
                        round(item.rating, 1) if item.rating else "", item.game_id]
                for col, val in enumerate(vals, 1):
                    c = ws.cell(row=row, column=col, value=val)
                    c.border = THIN_BORDER
                    if fill:
                        c.fill = fill
                    if col == 1:
                        c.alignment = CENTER
                row += 1

            row += 1  # 空行分隔

    def _write_summary_sheet(self, wb, groups: dict, fetch_date: str):
        ws = wb.create_sheet(title="汇总", index=0)
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 6

        # 二级表头：平台 × 榜单
        header_platforms = []
        for platform in PLATFORM_ORDER:
            for rank_type in RANK_TYPE_ORDER:
                if (platform, rank_type) in groups:
                    header_platforms.append((platform, rank_type))

        if not header_platforms:
            return

        # 行：1=平台名, 2=榜单名, 3..=排名数据
        ws.cell(row=1, column=1, value="排名")
        ws.cell(row=2, column=1, value="")

        for col_idx, (platform, rank_type) in enumerate(header_platforms, 2):
            c1 = ws.cell(row=1, column=col_idx, value=PLATFORMS.get(platform, platform))
            c1.font = HEADER_FONT; c1.fill = HEADER_FILL; c1.alignment = CENTER
            c2 = ws.cell(row=2, column=col_idx, value=RANK_TYPES.get(rank_type, rank_type))
            c2.font = Font(bold=True); c2.fill = PatternFill("solid", fgColor="BDD7EE")
            c2.alignment = CENTER
            ws.column_dimensions[get_column_letter(col_idx)].width = 28

        for rank_pos in range(1, 21):
            row = rank_pos + 2
            ws.cell(row=row, column=1, value=rank_pos).alignment = CENTER
            for col_idx, (platform, rank_type) in enumerate(header_platforms, 2):
                items = groups.get((platform, rank_type), [])
                matched = next((it for it in items if it.rank == rank_pos), None)
                ws.cell(row=row, column=col_idx, value=matched.name if matched else "—")
